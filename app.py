import os
import io
import csv
import json
import base64
import time
import logging
import aiohttp
import requests
from datetime import datetime, timedelta
from fastapi import FastAPI, Request, UploadFile, File, Header, HTTPException, Depends
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, Response
from fastapi.concurrency import run_in_threadpool
from fastapi.templating import Jinja2Templates
from fastapi.exceptions import RequestValidationError
from starlette.exceptions import HTTPException as StarletteHTTPException
from pydantic import BaseModel
from aiogram import Bot, Dispatcher
from aiogram.types import Update, ErrorEvent
from aiogram.fsm.storage.memory import MemoryStorage
from handlers import routers
from handlers.common import AccessMiddleware
from handlers.client_bot import create_client_dispatcher, set_commands as set_client_commands
from config import BOT_TOKEN, ADMIN_PASSWORD
from google_sheets import SheetManager, moscow_now
from bambu_cloud import BambuCloudManager

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import qrcode
from PIL import Image

from handlers_3mf import (
    extract_colors_from_3mf, hex_to_rgb, rgb_to_hex,
    group_similar_colors, find_closest_colors, generate_color_palette, BAMBU_COLORS,
)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

if not BOT_TOKEN:
    raise ValueError("BOT_TOKEN is required")

try:
    sheet_manager = SheetManager()
    logger.info("SheetManager успешно инициализирован")
except Exception as e:
    logger.error(f"Ошибка инициализации SheetManager: {e}")
    sheet_manager = None

# ---------- Основной бот ----------
bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())
dp.message.middleware(AccessMiddleware())
dp.callback_query.middleware(AccessMiddleware())
for router in routers:
    dp.include_router(router)


@dp.errors()
async def errors_handler(event: ErrorEvent):
    logger.error(f"Global error: {event.exception}", exc_info=event.exception)
    return True


# ---------- Клиентский бот ----------
CLIENT_BOT_TOKEN = os.getenv("CLIENT_BOT_TOKEN", "")
client_bot = Bot(token=CLIENT_BOT_TOKEN) if CLIENT_BOT_TOKEN else None
client_dp = create_client_dispatcher() if client_bot else None


app = FastAPI()
templates = Jinja2Templates(directory="templates")
bambu_cloud = BambuCloudManager()

PRICE_LOGO_URL = os.getenv("PRICE_LOGO_URL", "")
PUBLIC_URL = os.getenv("PUBLIC_URL", "")
PRICE_ORDER_CHAT_ID = os.getenv("PRICE_ORDER_CHAT_ID", "")
PRICE_ORDER_DEADLINE_DAYS = int(os.getenv("PRICE_ORDER_DEADLINE_DAYS", "7"))


# ---------- EXCEPTION HANDLERS ----------
@app.exception_handler(StarletteHTTPException)
async def http_exception_handler(request: Request, exc: StarletteHTTPException):
    if exc.status_code == 404:
        return templates.TemplateResponse("404.html", {"request": request}, status_code=404)
    return JSONResponse({"detail": exc.detail}, status_code=exc.status_code)


@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    return JSONResponse({"detail": "Неверный формат данных", "errors": exc.errors()}, status_code=422)


@app.exception_handler(Exception)
async def generic_exception_handler(request: Request, exc: Exception):
    logger.error(f"Unhandled exception at {request.url.path}: {exc}", exc_info=exc)
    try:
        return templates.TemplateResponse("500.html", {"request": request}, status_code=500)
    except Exception:
        return JSONResponse({"detail": "Internal Server Error"}, status_code=500)


# ---------- АВТОРИЗАЦИЯ ----------
def verify_admin(x_admin_token: str = Header(None, alias="X-Admin-Token")):
    if not ADMIN_PASSWORD:
        raise HTTPException(status_code=500, detail="ADMIN_PASSWORD не задан на сервере")
    if x_admin_token != ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Неверный пароль администратора")
    return True


class PriceItemIn(BaseModel):
    name: str
    description: str = ""
    photo: str = ""
    retail: str = ""
    wholesale: str = ""
    wholesale_from: str = ""
    category: str = ""


class LogoIn(BaseModel):
    url: str = ""


class BulkRowsIn(BaseModel):
    rows: list[int]


class BulkCategoryIn(BaseModel):
    rows: list[int]
    category: str


class OrderItemIn(BaseModel):
    name: str
    qty: int = 1


class OrderRequestIn(BaseModel):
    items: list[OrderItemIn]
    name: str = ""
    contact: str = ""
    comment: str = ""
    promo_code: str = ""


class PromoIn(BaseModel):
    code: str


class ReviewIn(BaseModel):
    product_row: int
    name: str = ""
    rating: int = 5
    text: str


# ---------- ЛОГОТИП ----------
def _get_price_logo_url() -> str:
    if sheet_manager:
        try:
            v = sheet_manager.get_setting("price_logo_url", "")
            if v:
                return v
        except Exception as e:
            logger.warning(f"get_setting(price_logo_url) failed: {e}")
    return PRICE_LOGO_URL


@app.get("/api/price/logo")
async def get_price_logo(_: bool = Depends(verify_admin)):
    return JSONResponse({"url": _get_price_logo_url()})


@app.post("/api/price/logo")
async def set_price_logo(payload: LogoIn, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    if not payload.url:
        return JSONResponse({"error": "URL не может быть пустым"}, status_code=400)
    ok = sheet_manager.set_setting("price_logo_url", payload.url)
    if not ok:
        return JSONResponse({"error": "Не удалось сохранить логотип"}, status_code=500)
    return JSONResponse({"status": "ok", "url": payload.url})


@app.delete("/api/price/logo")
async def delete_price_logo(_: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    sheet_manager.delete_setting("price_logo_url")
    return JSONResponse({"status": "ok"})


# ---------- ЗАГРУЗКА ФОТО ----------
DRIVE_UPLOAD_URL = os.getenv("DRIVE_UPLOAD_URL", "")
DRIVE_UPLOAD_SECRET = os.getenv("DRIVE_UPLOAD_SECRET", "")
MAX_IMAGE_SIZE = 5 * 1024 * 1024


@app.post("/api/upload_image")
async def upload_image_api(file: UploadFile = File(...), _: bool = Depends(verify_admin)):
    if not DRIVE_UPLOAD_URL or not DRIVE_UPLOAD_SECRET:
        return JSONResponse({"error": "DRIVE_UPLOAD_URL / DRIVE_UPLOAD_SECRET не заданы"}, status_code=500)
    if not file.content_type or not file.content_type.startswith("image/"):
        return JSONResponse({"error": "Только изображения"}, status_code=400)
    content = await file.read()
    if not content:
        return JSONResponse({"error": "Пустой файл"}, status_code=400)
    if len(content) > MAX_IMAGE_SIZE:
        return JSONResponse({"error": "Файл больше 5 МБ"}, status_code=400)

    b64 = base64.b64encode(content).decode("ascii")
    filename = file.filename or f"upload_{int(time.time())}.jpg"
    payload = {"secret": DRIVE_UPLOAD_SECRET, "filename": filename, "mime": file.content_type, "image": b64}
    try:
        async with aiohttp.ClientSession() as session:
            async with session.post(DRIVE_UPLOAD_URL, json=payload,
                                    timeout=aiohttp.ClientTimeout(total=60)) as resp:
                raw = await resp.text()
                if resp.status != 200:
                    logger.error(f"Drive upload HTTP {resp.status}: {raw[:300]}")
                    return JSONResponse({"error": "Google Drive вернул ошибку"}, status_code=502)
                data = json.loads(raw)
    except Exception as e:
        logger.error(f"Drive upload exception: {e}")
        return JSONResponse({"error": f"Ошибка соединения с Drive: {e}"}, status_code=502)
    if not data.get("ok"):
        return JSONResponse({"error": data.get("error") or "Drive отказал"}, status_code=502)
    return JSONResponse({"url": data["url"]})


# ---------- QR ----------
def _get_public_base_url(request: Request) -> str:
    if PUBLIC_URL:
        return PUBLIC_URL.rstrip("/")
    proto = request.headers.get("x-forwarded-proto") or request.url.scheme
    host = request.headers.get("x-forwarded-host") or request.headers.get("host") or ""
    if host:
        return f"{proto}://{host}"
    return str(request.base_url).rstrip("/")


@app.get("/api/price/qrcode")
async def price_qrcode(request: Request, _: bool = Depends(verify_admin)):
    try:
        base = _get_public_base_url(request)
        url = f"{base}/price"
        img = qrcode.make(url, box_size=10, border=2)
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        return Response(content=buf.getvalue(), media_type="image/png",
                        headers={"Cache-Control": "no-store"})
    except Exception as e:
        logger.error(f"QR error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


# ---------- WORD ----------
def _flatten_on_white(img):
    if img.mode == 'P':
        img = img.convert('RGBA')
    if img.mode in ('RGBA', 'LA'):
        bg = Image.new('RGB', img.size, (255, 255, 255))
        alpha = img.split()[-1]
        bg.paste(img, mask=alpha)
        return bg
    if img.mode != 'RGB':
        return img.convert('RGB')
    return img


def _download_image(url: str, timeout: int = 10):
    try:
        resp = requests.get(url, timeout=timeout, stream=True, allow_redirects=True)
        if resp.status_code != 200:
            return None
        ct = (resp.headers.get("Content-Type") or "").lower()
        if not ct.startswith("image/"):
            return None
        data = resp.content
        if not data or len(data) > MAX_IMAGE_SIZE:
            return None
        try:
            img = Image.open(io.BytesIO(data))
            img = _flatten_on_white(img)
            out = io.BytesIO()
            img.save(out, format='PNG', optimize=True)
            out.seek(0)
            return out
        except Exception:
            return io.BytesIO(data)
    except Exception as e:
        logger.warning(f"download {url}: {e}")
        return None


def _add_page_number_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, end):
        run._r.append(el)


def _setup_header_footer(doc, logo_url):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    try:
        first_header = section.first_page_header
        first_header.is_linked_to_previous = False
        p = first_header.paragraphs[0] if first_header.paragraphs else first_header.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if logo_url:
            logo_io = _download_image(logo_url)
            if logo_io:
                try:
                    p.add_run().add_picture(logo_io, width=Cm(3.0))
                except Exception:
                    pass
    except Exception:
        pass
    try:
        regular_header = section.header
        regular_header.is_linked_to_previous = False
        for para in regular_header.paragraphs:
            para.text = ""
    except Exception:
        pass
    try:
        footer = section.footer
        footer.is_linked_to_previous = False
        ftr_p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        ftr_p.text = ""
        ftr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = ftr_p.add_run(f"Прайс-лист · сгенерировано {datetime.now().strftime('%d.%m.%Y %H:%M')} · стр. ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        run.italic = True
        _add_page_number_field(ftr_p)
        for r in ftr_p.runs[1:]:
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            r.italic = True
    except Exception:
        pass
    try:
        section.first_page_footer.is_linked_to_previous = True
    except Exception:
        pass


def _build_price_docx(items):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    _setup_header_footer(doc, _get_price_logo_url())
    title = doc.add_heading("Прайс-лист", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not items:
        doc.add_paragraph("Прайс пуст.")
        buf = io.BytesIO(); doc.save(buf); buf.seek(0)
        return buf.getvalue()
    by_cat = {}
    for it in items:
        cat = (it.get("category") or "Без категории").strip() or "Без категории"
        by_cat.setdefault(cat, []).append(it)
    sorted_cats = sorted(by_cat.keys(), key=lambda s: s.lower())
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = toc_title.add_run("Содержание")
    r.bold = True
    r.font.size = Pt(14)
    for cat in sorted_cats:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(cat)
        run.font.size = Pt(11)
        cnt_run = p.add_run(f"  —  {len(by_cat[cat])} поз.")
        cnt_run.font.size = Pt(9)
        cnt_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        cnt_run.italic = True
    doc.add_page_break()
    for cat in sorted_cats:
        doc.add_heading(cat, level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = table._tbl.tblPr
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        tblPr.append(layout)
        hdr = table.rows[0].cells
        for i, h in enumerate(["Фото", "Название", "Описание", "Розница", "Опт"]):
            hdr[i].text = ""
            pp = hdr[i].paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = pp.add_run(h)
            run.bold = True
            run.font.size = Pt(10)
        widths = [Cm(2.2), Cm(3.5), Cm(5.5), Cm(2.2), Cm(2.5)]
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
        for it in by_cat[cat]:
            row = table.add_row().cells
            for i, w in enumerate(widths):
                row[i].width = w
                row[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            photo_url = (it.get("photo") or "").strip()
            p_photo = row[0].paragraphs[0]
            p_photo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if photo_url:
                img_io = _download_image(photo_url)
                if img_io:
                    try:
                        p_photo.add_run().add_picture(img_io, width=Cm(2.0))
                    except Exception:
                        p_photo.add_run("—")
                else:
                    p_photo.add_run("—")
            else:
                p_photo.add_run("—")
            name_p = row[1].paragraphs[0]
            name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            name_p.add_run(it.get("name", "") or "")
            desc_p = row[2].paragraphs[0]
            desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            desc_p.add_run((it.get("description", "") or "").replace("\r\n", "\n"))
            retail = (it.get("retail") or "").strip()
            wholesale = (it.get("wholesale") or "").strip()
            wholesale_from = (it.get("wholesale_from") or "").strip()
            retail_p = row[3].paragraphs[0]
            retail_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if retail:
                rr = retail_p.add_run(f"{retail} ₽"); rr.bold = True
            else:
                retail_p.add_run("—")
            ws_p = row[4].paragraphs[0]
            ws_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if wholesale:
                ws_p.add_run(f"{wholesale} ₽")
                if wholesale_from:
                    fp = row[4].add_paragraph(f"от {wholesale_from} шт")
                    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fp.runs[0].font.size = Pt(9)
                    fp.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            else:
                ws_p.add_run("—")
        doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = summary.add_run(f"Всего товаров: {len(items)} · Категорий: {len(by_cat)}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return buf.getvalue()


@app.get("/api/price/export_docx")
async def export_price_docx(_: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        items = sheet_manager.get_price_items()
        content = await run_in_threadpool(_build_price_docx, items)
        filename = f"price_{datetime.now().strftime('%Y-%m-%d')}.docx"
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except Exception as e:
        logger.error(f"Ошибка экспорта в Word: {e}", exc_info=e)
        return JSONResponse({"error": str(e)}, status_code=500)


def get_days_left(deadline):
    try:
        due = datetime.strptime(deadline, "%Y-%m-%d")
        diff = (due - datetime.now()).days
        if diff < 0: return "Просрочено"
        if diff == 0: return "Сегодня"
        if diff == 1: return "Завтра"
        return f"{diff} дн."
    except:
        return "—"


# ---------- WEBHOOKS ----------
@app.post("/webhook")
async def webhook(request: Request):
    try:
        update = Update(**(await request.json()))
        await dp.feed_update(bot, update)
        return {"status": "ok"}
    except Exception as e:
        logger.error(f"Webhook error: {e}")
        return {"status": "error"}


@app.post("/webhook_client")
async def webhook_client(request: Request):
    if not client_bot or not client_dp:
        return JSONResponse({"error": "Client bot not configured"}, status_code=503)
    try:
        update = Update(**(await request.json()))
        await client_dp.feed_update(client_bot, update)
        return {"status": "ok"}
    except Exception as e:
        logger.error(f"Client webhook error: {e}")
        return {"status": "error"}


# ---------- СТРАНИЦЫ ----------
@app.get("/")
async def root():
    return {"status": "3D Printer Bot is running"}


@app.get("/tracker", response_class=HTMLResponse)
async def tracker_page(request: Request):
    return templates.TemplateResponse("tracker.html", {"request": request})


@app.get("/price", response_class=HTMLResponse)
async def price_page(request: Request):
    return templates.TemplateResponse("price.html", {"request": request})


@app.get("/upload_3mf", response_class=HTMLResponse)
async def upload_3mf_page():
    return FileResponse("upload_3mf.html")


@app.get("/share/{row}", response_class=HTMLResponse)
async def share_page(request: Request, row: int):
    if not sheet_manager:
        return templates.TemplateResponse("404.html", {"request": request}, status_code=404)
    it = sheet_manager.get_price_item_by_row(row)
    if not it:
        return templates.TemplateResponse("404.html", {"request": request}, status_code=404)
    return templates.TemplateResponse("share.html", {
        "request": request,
        "item": it,
        "public_url": _get_public_base_url(request),
    })


@app.get("/sitemap.xml")
async def sitemap(request: Request):
    base = _get_public_base_url(request)
    urls = [f"{base}/", f"{base}/price"]
    if sheet_manager:
        try:
            for it in sheet_manager.get_price_items():
                urls.append(f"{base}/share/{it['row_index']}")
        except Exception as e:
            logger.warning(f"sitemap: {e}")
    xml = '<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
    for u in urls:
        xml += f"  <url><loc>{u}</loc></url>\n"
    xml += "</urlset>"
    return Response(content=xml, media_type="application/xml")


@app.get("/robots.txt")
async def robots(request: Request):
    base = _get_public_base_url(request)
    text = (
        "User-agent: *\n"
        "Allow: /\n"
        "Disallow: /tracker\n"
        "Disallow: /api/\n"
        "Disallow: /webhook\n"
        "Disallow: /webhook_client\n"
        "Disallow: /check_tasks\n"
        f"Sitemap: {base}/sitemap.xml\n"
    )
    return Response(content=text, media_type="text/plain")


@app.get("/manifest.json")
async def manifest():
    return JSONResponse(content={
        "name": "Kobra Print — прайс",
        "short_name": "Kobra Print",
        "start_url": "/price",
        "display": "standalone",
        "background_color": "#ffffff",
        "theme_color": "#3b82f6",
        "icons": [{"src": "/static/icon-192.png", "sizes": "192x192", "type": "image/png"}]
    })


# ---------- API: ЗАКАЗЫ ----------
@app.get("/api/orders")
async def get_orders_api(customer: str = "", from_date: str = "", to_date: str = ""):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        result = []
        for order in sheet_manager.get_active_orders():
            if len(order) < 8:
                continue
            order_customer = order[7] if order[7] else ""
            if customer and customer.lower() not in order_customer.lower():
                continue
            if from_date and order[4] < from_date:
                continue
            if to_date and order[4] > to_date:
                continue
            try:
                ordered = int(order[2]); printed = int(order[3])
            except (ValueError, TypeError):
                ordered = 0; printed = 0
            result.append({
                "id": order[0], "position": order[1],
                "ordered": ordered, "printed": printed,
                "deadline": order[4], "modified": order[5],
                "status": order[6], "customer": order_customer,
                "progress": round(printed / ordered * 100) if ordered > 0 else 0
            })
        return JSONResponse({"orders": result})
    except Exception as e:
        logger.error(f"API error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


# ---------- API: ЗАДАЧИ ----------
@app.get("/api/tasks")
async def get_tasks_api():
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        result = [{
            "id": t[0], "title": t[1], "deadline": t[2],
            "time": t[3] if len(t) > 3 else "",
            "assignee": t[4] if t[4] else "Общая",
            "status": t[5],
            "time_left": get_days_left(t[2])
        } for t in sheet_manager.get_active_tasks()]
        return JSONResponse({"tasks": result})
    except Exception as e:
        logger.error(f"API tasks error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


# ---------- API: ПРИНТЕРЫ ----------
@app.get("/api/printers")
async def get_printers_api():
    printers = await run_in_threadpool(bambu_cloud.get_printers)
    return JSONResponse({"printers": printers})


# ---------- API: 3MF ----------
MAX_3MF_SIZE = 50 * 1024 * 1024


def _process_3mf_bytes(content: bytes):
    raw_colors_hex = extract_colors_from_3mf(content)
    if not raw_colors_hex:
        return {"colors": [], "count": 0, "palette": None}
    raw_colors_rgb = [hex_to_rgb(h) for h in raw_colors_hex]
    grouped = group_similar_colors(raw_colors_rgb, tolerance=20, max_colors=10)
    unique_matches = []
    seen_hex = set()
    for rgb in grouped:
        matches = find_closest_colors(rgb_to_hex(rgb), BAMBU_COLORS, top_n=1)
        if matches and matches[0][1]["hex"] not in seen_hex:
            seen_hex.add(matches[0][1]["hex"])
            unique_matches.append(matches[0][1])
    palette_b64 = None
    try:
        palette_bytes = generate_color_palette(grouped)
        if palette_bytes:
            palette_b64 = base64.b64encode(palette_bytes).decode("ascii")
    except Exception as e:
        logger.warning(f"palette error: {e}")
    return {"colors": unique_matches, "count": len(unique_matches), "palette": palette_b64}


@app.post("/api/analyze_3mf")
async def analyze_3mf_api(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".3mf"):
        return JSONResponse({"detail": "Нужен файл .3mf"}, status_code=400)
    try:
        content = await file.read()
    except Exception as e:
        return JSONResponse({"detail": f"Не удалось прочитать файл: {e}"}, status_code=400)
    if len(content) > MAX_3MF_SIZE:
        return JSONResponse({"detail": "Файл больше 50 МБ"}, status_code=400)
    try:
        result = await run_in_threadpool(_process_3mf_bytes, content)
        return JSONResponse(result)
    except Exception as e:
        logger.error(f"3MF error: {e}", exc_info=e)
        return JSONResponse({"detail": f"Ошибка анализа: {e}"}, status_code=500)


# ---------- API: ПРАЙС ----------
@app.get("/api/price")
async def get_price_api():
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        items = sheet_manager.get_price_items()
        categories = sheet_manager.get_price_categories()
        public_items = [{
            "row_index": it["row_index"],
            "name": it["name"], "description": it["description"],
            "photo": it["photo"], "retail": it["retail"],
            "wholesale": it["wholesale"], "wholesale_from": it["wholesale_from"],
            "category": it["category"],
        } for it in items]
        return JSONResponse({"items": public_items, "categories": categories})
    except Exception as e:
        logger.error(f"API price error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


@app.get("/api/price/admin")
async def get_price_admin_api(_: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        return JSONResponse({"items": sheet_manager.get_price_items()})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/api/price")
async def add_price_api(payload: PriceItemIn, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        ok = sheet_manager.add_price_item(
            payload.name, payload.description, payload.photo,
            payload.retail, payload.wholesale, payload.wholesale_from, payload.category
        )
        if not ok:
            return JSONResponse({"error": "Не удалось добавить"}, status_code=400)
        return JSONResponse({"status": "ok"})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.put("/api/price/{row_index}")
async def update_price_api(row_index: int, payload: PriceItemIn, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    if not sheet_manager.get_price_item_by_row(row_index):
        return JSONResponse({"error": "Товар не найден"}, status_code=404)
    try:
        for field, value in [
            ("name", payload.name), ("description", payload.description),
            ("photo", payload.photo), ("retail", payload.retail),
            ("wholesale", payload.wholesale), ("wholesale_from", payload.wholesale_from),
            ("category", payload.category),
        ]:
            sheet_manager.update_price_item(row_index, field, value)
        return JSONResponse({"status": "ok"})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.delete("/api/price/{row_index}")
async def delete_price_api(row_index: int, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    if not sheet_manager.get_price_item_by_row(row_index):
        return JSONResponse({"error": "Товар не найден"}, status_code=404)
    try:
        ok = sheet_manager.delete_price_item(row_index)
        if not ok:
            return JSONResponse({"error": "Не удалось удалить"}, status_code=400)
        return JSONResponse({"status": "ok"})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


# ---------- API: МАССОВЫЕ ДЕЙСТВИЯ ----------
@app.post("/api/price/bulk_delete")
async def bulk_delete_api(payload: BulkRowsIn, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    if not payload.rows:
        return JSONResponse({"error": "Пустой список"}, status_code=400)
    try:
        deleted = sheet_manager.delete_price_items(payload.rows)
        return JSONResponse({"status": "ok", "deleted": deleted})
    except Exception as e:
        logger.error(f"bulk_delete error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/api/price/bulk_set_category")
async def bulk_set_category_api(payload: BulkCategoryIn, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    if not payload.rows:
        return JSONResponse({"error": "Пустой список"}, status_code=400)
    if not payload.category.strip():
        return JSONResponse({"error": "Категория не может быть пустой"}, status_code=400)
    try:
        updated = sheet_manager.set_category_for_rows(payload.rows, payload.category.strip())
        return JSONResponse({"status": "ok", "updated": updated})
    except Exception as e:
        logger.error(f"bulk_set_category error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


# ---------- API: CSV ----------
@app.post("/api/price/import_csv")
async def import_csv_api(file: UploadFile = File(...), _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    if not file.filename or not file.filename.lower().endswith(".csv"):
        return JSONResponse({"error": "Нужен файл .csv"}, status_code=400)
    try:
        content = await file.read()
    except Exception as e:
        return JSONResponse({"error": f"Не удалось прочитать файл: {e}"}, status_code=400)
    if not content:
        return JSONResponse({"error": "Пустой файл"}, status_code=400)
    if len(content) > 5 * 1024 * 1024:
        return JSONResponse({"error": "Файл больше 5 МБ"}, status_code=400)

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1251")
        except Exception:
            return JSONResponse({"error": "Не удалось декодировать CSV"}, status_code=400)

    try:
        reader = csv.DictReader(io.StringIO(text))
        raw_items = []
        for row in reader:
            name = (row.get("name") or row.get("Название") or "").strip()
            if not name:
                continue
            raw_items.append({
                "name": name,
                "description": (row.get("description") or row.get("Описание") or "").strip(),
                "photo": (row.get("photo") or row.get("Фото") or "").strip(),
                "retail": (row.get("retail") or row.get("Розница") or "").strip(),
                "wholesale": (row.get("wholesale") or row.get("Опт") or "").strip(),
                "wholesale_from": (row.get("wholesale_from") or row.get("Опт от") or "").strip(),
                "category": (row.get("category") or row.get("Категория") or "Без категории").strip() or "Без категории",
            })
    except Exception as e:
        return JSONResponse({"error": f"Ошибка парсинга CSV: {e}"}, status_code=400)

    if not raw_items:
        return JSONResponse({"error": "В CSV не найдено ни одной строки"}, status_code=400)

    try:
        added = await run_in_threadpool(sheet_manager.import_price_items, raw_items)
        return JSONResponse({"status": "ok", "added": added})
    except Exception as e:
        logger.error(f"import_csv error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


# ---------- API: ПРОМОКОДЫ ----------
@app.post("/api/price/validate_promo")
async def validate_promo_api(payload: PromoIn):
    if not sheet_manager:
        return JSONResponse({"valid": False, "reason": "Сервис недоступен"}, status_code=500)
    result = sheet_manager.validate_promo(payload.code)
    return JSONResponse(result)


# ---------- API: ОТЗЫВЫ (с модерацией) ----------
# ⚠️ ВАЖНО: конкретные пути идут ПЕРЕД /{row}, иначе "pending" попытается стать int → 422
@app.get("/api/price/reviews_summary")
async def get_reviews_summary_api():
    if not sheet_manager:
        return JSONResponse({})
    try:
        return JSONResponse(sheet_manager.get_reviews_summary())
    except Exception as e:
        logger.error(f"reviews_summary error: {e}")
        return JSONResponse({})


@app.get("/api/price/reviews/pending")
async def get_pending_reviews_api(_: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"items": []})
    try:
        items = sheet_manager.get_pending_reviews()
        enriched = []
        for r in items:
            product = sheet_manager.get_price_item_by_row(r["product_row"])
            r_copy = dict(r)
            r_copy["product_name"] = product["name"] if product else f"Товар #{r['product_row']}"
            enriched.append(r_copy)
        return JSONResponse({"items": enriched})
    except Exception as e:
        logger.error(f"pending_reviews error: {e}")
        return JSONResponse({"items": []})


@app.get("/api/price/reviews/{row}")
async def get_reviews_api(row: int):
    """Публичный — только одобренные. ⚠️ Объявлен ПОСЛЕ конкретных путей."""
    if not sheet_manager:
        return JSONResponse({"items": []})
    try:
        items = sheet_manager.get_reviews_for_product(row, limit=30)
        return JSONResponse({"items": items})
    except Exception as e:
        logger.error(f"get_reviews error: {e}")
        return JSONResponse({"items": []})


@app.post("/api/price/review/{review_id}/approve")
async def approve_review_api(review_id: int, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    ok = sheet_manager.approve_review(review_id)
    if not ok:
        return JSONResponse({"error": "Не удалось одобрить"}, status_code=400)
    return JSONResponse({"status": "ok"})


@app.post("/api/price/review/{review_id}/reject")
async def reject_review_api(review_id: int, _: bool = Depends(verify_admin)):
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    ok = sheet_manager.reject_review(review_id)
    if not ok:
        return JSONResponse({"error": "Не удалось отклонить"}, status_code=400)
    return JSONResponse({"status": "ok"})


@app.post("/api/price/review")
async def add_review_api(payload: ReviewIn):
    if not sheet_manager:
        return JSONResponse({"error": "Сервис недоступен"}, status_code=500)
    product = sheet_manager.get_price_item_by_row(payload.product_row)
    if not product:
        return JSONResponse({"error": "Товар не найден"}, status_code=404)
    if payload.rating < 1 or payload.rating > 5:
        return JSONResponse({"error": "Оценка 1–5"}, status_code=400)
    if not payload.text.strip():
        return JSONResponse({"error": "Введите текст отзыва"}, status_code=400)
    ok = sheet_manager.add_review(payload.product_row, payload.name, payload.rating, payload.text)
    if not ok:
        return JSONResponse({"error": "Не удалось сохранить отзыв"}, status_code=500)

    try:
        recipients = []
        if PRICE_ORDER_CHAT_ID:
            try:
                recipients = [int(PRICE_ORDER_CHAT_ID)]
            except ValueError:
                pass
        if not recipients:
            try:
                from config import ALLOWED_USERS
                recipients = list(ALLOWED_USERS or [])
            except Exception:
                pass

        stars = "★" * payload.rating + "☆" * (5 - payload.rating)
        product_name = product.get("name", "—")
        msg = (
            f"📝 <b>Новый отзыв на модерации</b>\n\n"
            f"Товар: {product_name}\n"
            f"Оценка: {stars} ({payload.rating}/5)\n"
            f"Имя: {payload.name or 'Аноним'}\n\n"
            f"Текст:\n{payload.text[:300]}"
        )
        for rid in recipients:
            try:
                await bot.send_message(rid, msg, parse_mode="HTML")
            except Exception as e:
                logger.warning(f"review notify to {rid}: {e}")
    except Exception as e:
        logger.error(f"review notify error: {e}")

    return JSONResponse({"status": "ok", "moderation": True})


# ---------- API: ЗАЯВКА С ПРАЙСА ----------
@app.post("/api/price/order_request")
async def price_order_request(payload: OrderRequestIn):
    if not payload.items:
        return JSONResponse({"error": "Корзина пуста"}, status_code=400)
    if not payload.contact.strip():
        return JSONResponse({"error": "Укажите контакт для связи"}, status_code=400)

    customer_name = payload.name.strip() or "Клиент с прайса"
    customer_full = f"{customer_name} ({payload.contact.strip()})"
    deadline_str = (datetime.now() + timedelta(days=PRICE_ORDER_DEADLINE_DAYS)).strftime("%Y-%m-%d")

    promo_info = None
    promo_code = (payload.promo_code or "").strip()
    if promo_code and sheet_manager:
        promo_info = sheet_manager.validate_promo(promo_code)
        if not promo_info.get("valid"):
            promo_info = None

    created_orders = []
    if sheet_manager:
        for it in payload.items:
            try:
                position = f"С прайса: {it.name}"
                num = sheet_manager.add_order(position, it.qty, deadline_str, customer_full)
                created_orders.append(num)
            except Exception as e:
                logger.error(f"Не удалось создать заказ для '{it.name}': {e}")

    if promo_info:
        sheet_manager.increment_promo_usage(promo_info["code"])

    lines = ["🛒 <b>Новая заявка с прайса</b>", ""]
    if payload.name.strip():
        lines.append(f"👤 Имя: {payload.name.strip()}")
    lines.append(f"📞 Контакт: {payload.contact.strip()}")
    if payload.comment.strip():
        lines.append(f"💬 Комментарий: {payload.comment.strip()}")
    if promo_info:
        lines.append(f"🎁 Промокод: <b>{promo_info['code']}</b> (-{promo_info['discount_percent']}%)")

    if created_orders:
        nums = ", ".join(f"№{n}" for n in created_orders)
        lines.append(f"📦 <b>Созданы заказы:</b> {nums}")
        lines.append(f"📅 Срок (по умолчанию): {deadline_str}")

    lines.append("")
    lines.append("📋 Состав:")
    for it in payload.items:
        lines.append(f"• {it.name} — {it.qty} шт")
    lines.append("")
    lines.append(f"Всего позиций: {len(payload.items)}")
    text = "\n".join(lines)

    recipients = []
    if PRICE_ORDER_CHAT_ID:
        try:
            recipients = [int(PRICE_ORDER_CHAT_ID)]
        except ValueError:
            recipients = []
    if not recipients:
        try:
            from config import ALLOWED_USERS
            recipients = list(ALLOWED_USERS or [])
        except Exception:
            recipients = []

    if not recipients:
        return JSONResponse({"error": "Получатель не настроен (PRICE_ORDER_CHAT_ID)"}, status_code=500)

    sent = 0
    for rid in recipients:
        try:
            await bot.send_message(rid, text, parse_mode="HTML")
            sent += 1
        except Exception as e:
            logger.error(f"order_request send to {rid}: {e}")

    return JSONResponse({
        "status": "ok",
        "sent": sent,
        "orders": created_orders,
        "deadline": deadline_str,
        "promo_applied": bool(promo_info),
    })


# ---------- CRON ----------
@app.get("/check_tasks")
async def check_tasks():
    if not sheet_manager:
        return JSONResponse({"error": "SheetManager не инициализирован"}, status_code=500)
    try:
        tasks = sheet_manager.get_tasks_for_notification()
        now = moscow_now()
        notified_count = 0
        for task in tasks:
            try:
                diff_minutes = (task["deadline_dt"] - now).total_seconds() / 60
                assignee = task["assignee"]
                task_id = task["id"]
                title = task["title"]
                user_settings = sheet_manager.get_user_settings(assignee) if assignee else None
                morning_hour = int(user_settings.get("morning_time", "09:00").split(':')[0]) if user_settings else 9
                recipients = [int(assignee)] if (assignee and str(assignee).isdigit()) else sheet_manager.get_all_subscribers()

                if now.hour == morning_hour and now.minute == 0 and task["notified_morning"] == "0":
                    for recipient in recipients:
                        try:
                            await bot.send_message(recipient, f"🌅 Напоминание: сегодня задача '{title}' должна быть выполнена до {task['deadline_dt'].strftime('%H:%M')}!")
                            notified_count += 1
                        except Exception as e:
                            logger.error(f"morning notify: {e}")
                    sheet_manager.update_task_notification(task_id, 'notified_morning', '1')

                for minutes, field in [(60, 'notified_60'), (30, 'notified_30'), (15, 'notified_15'), (0, 'notified_0')]:
                    if abs(diff_minutes - minutes) < 0.5 and task[field] == "0":
                        for recipient in recipients:
                            try:
                                text = (f"🔔 Срок задачи '{title}' истёк (до {task['deadline_dt'].strftime('%H:%M')})!"
                                        if minutes == 0
                                        else f"⏰ Через {minutes} минут задача '{title}' должна быть выполнена (до {task['deadline_dt'].strftime('%H:%M')})!")
                                await bot.send_message(recipient, text)
                                notified_count += 1
                            except Exception as e:
                                logger.error(f"notify {minutes}: {e}")
                        sheet_manager.update_task_notification(task_id, field, '1')
            except Exception as e:
                logger.error(f"task {task.get('id')}: {e}")
        return JSONResponse({"status": "ok", "notified": notified_count})
    except Exception as e:
        logger.error(f"check_tasks: {e}")
        return JSONResponse({"status": "error", "message": str(e)}, status_code=500)


@app.on_event("startup")
async def on_startup():
    if client_bot:
        try:
            await set_client_commands(client_bot)
            logger.info("Client bot commands registered")
        except Exception as e:
            logger.warning(f"client bot setup: {e}")


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
